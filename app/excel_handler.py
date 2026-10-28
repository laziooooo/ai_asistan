"""Excel Dosya İşleme Modülü"""
import os
import csv
import json
import logging
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Optional

logger = logging.getLogger(__name__)

try:
    import openpyxl
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logger.warning("openpyxl kütüphanesi bulunamadı. pip install openpyxl")

from .config import config, DATA_DIR

class ExcelHandler:
    """Excel dosya işlemleri"""
    
    def __init__(self):
        self.workbook = None
        self.worksheet = None
        self.file_path = None
        self.headers = ['Zaman', 'Gönderen', 'Sütun', 'Değer', 'Mesaj', 'Durum']
        self.records = []
        self._load_or_create()
    
    def _load_or_create(self):
        """Excel dosyasını yükle veya oluştur"""
        if not OPENPYXL_AVAILABLE:
            logger.error("openpyxl yok, CSV moduna geçiliyor")
            self._init_csv()
            return
        
        try:
            # Varsayılan dosya yolu
            default_path = DATA_DIR / 'asistan_kayitlar.xlsx'
            
            if default_path.exists():
                self.file_path = default_path
                self.workbook = load_workbook(str(default_path))
                self.worksheet = self.workbook.active
                
                # Başlıkları oku
                if self.worksheet.max_row > 1:
                    self.headers = [cell.value for cell in self.worksheet[1]]
                    logger.info(f"Mevcut Excel yüklendi: {len(self.worksheet.max_row - 1)} kayıt")
            else:
                self.file_path = default_path
                self.workbook = Workbook()
                self.worksheet = self.workbook.active
                self.worksheet.title = "Kayıtlar"
                self._create_headers()
                self.workbook.save(str(default_path))
                logger.info("Yeni Excel dosyası oluşturuldu")
            
        except Exception as e:
            logger.error(f"Excel yükleme hatası: {e}")
            self._init_csv()
    
    def _init_csv(self):
        """CSV modunu başlat"""
        self.csv_path = DATA_DIR / 'asistan_kayitlar.csv'
        if self.csv_path.exists():
            with open(self.csv_path, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                self.records = list(reader)
        else:
            with open(self.csv_path, 'w', encoding='utf-8', newline='') as f:
                writer = csv.DictWriter(f, fieldnames=self.headers)
                writer.writeheader()
    
    def _create_headers(self):
        """Başlıkları oluştur"""
        if self.worksheet:
            for col, header in enumerate(self.headers, 1):
                cell = self.worksheet.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            
            # Sütun genişliklerini ayarla
            self.worksheet.column_dimensions['A'].width = 20  # Zaman
            self.worksheet.column_dimensions['B'].width = 15  # Gönderen
            self.worksheet.column_dimensions['C'].width = 15  # Sütun
            self.worksheet.column_dimensions['D'].width = 30  # Değer
            self.worksheet.column_dimensions['E'].width = 50  # Mesaj
            self.worksheet.column_dimensions['F'].width = 10  # Durum
    
    def add_record(self, record: Dict):
        """Yeni kayıt ekle"""
        try:
            # Durum ekle
            record['Durum'] = 'Yeni'
            record['Zaman'] = datetime.now().strftime('%d.%m.%Y %H:%M:%S')
            
            if OPENPYXL_AVAILABLE and self.worksheet:
                # Excel'e ekle
                row = self.worksheet.max_row + 1
                for col, header in enumerate(self.headers, 1):
                    value = record.get(header, '')
                    self.worksheet.cell(row=row, column=col, value=str(value))
                
                self.workbook.save(str(self.file_path))
                logger.info(f"Excel'e kayıt eklendi: satır {row}")
            else:
                # CSV'ye ekle
                self.records.append(record)
                with open(self.csv_path, 'a', encoding='utf-8', newline='') as f:
                    writer = csv.DictWriter(f, fieldnames=self.headers)
                    writer.writerow(record)
                logger.info(f"CSV'ye kayıt eklendi")
            
            return True
            
        except Exception as e:
            logger.error(f"Kayıt ekleme hatası: {e}")
            return False
    
    def get_all_records(self) -> List[Dict]:
        """Tüm kayıtları getir"""
        try:
            if OPENPYXL_AVAILABLE and self.worksheet:
                records = []
                for row in self.worksheet.iter_rows(min_row=2, values_only=True):
                    if row[0]:  # Zaman sütunu boş değilse
                        record = {}
                        for col, header in enumerate(self.headers):
                            record[header] = row[col] if col < len(row) else ''
                        records.append(record)
                return records
            else:
                return self.records
        except Exception as e:
            logger.error(f"Kayıt okuma hatası: {e}")
            return []
    
    def update_record_status(self, index: int, status: str):
        """Kayıt durumunu güncelle"""
        try:
            if OPENPYXL_AVAILABLE and self.worksheet:
                row = index + 2  # +1 başlık, +1 index
                status_col = self.headers.index('Durum') + 1
                self.worksheet.cell(row=row, column=status_col, value=status)
                self.workbook.save(str(self.file_path))
                logger.info(f"Kayıt durumu güncellendi: {status}")
        except Exception as e:
            logger.error(f"Durum güncelleme hatası: {e}")
    
    def get_pending_records(self) -> List[Dict]:
        """Bekleyen kayıtları getir (Durum = Yeni)"""
        all_records = self.get_all_records()
        return [r for r in all_records if r.get('Durum') == 'Yeni']
    
    def get_high_priority_records(self) -> List[Dict]:
        """Yüksek öncelikli kayıtları getir"""
        all_records = self.get_all_records()
        high_priority = []
        for r in all_records:
            mesaj = str(r.get('Mesaj', '')).lower()
            if any(w in mesaj for w in ['acil', 'hemen', 'öncelikli', 'important', 'urgent']):
                high_priority.append(r)
        return high_priority
    
    def export_to_csv(self, output_path: str = None):
        """Excel'i CSV'ye aktar"""
        try:
            if not output_path:
                output_path = DATA_DIR / 'export.csv'
            
            records = self.get_all_records()
            with open(output_path, 'w', encoding='utf-8', newline='') as f:
                writer = csv.DictWriter(f, fieldnames=self.headers)
                writer.writeheader()
                writer.writerows(records)
            
            logger.info(f"CSV'ye aktarıldı: {output_path}")
            return True
            
        except Exception as e:
            logger.error(f"CSV aktarma hatası: {e}")
            return False
    
    def get_stats(self) -> Dict:
        """İstatistikleri getir"""
        all_records = self.get_all_records()
        pending = self.get_pending_records()
        high_priority = self.get_high_priority_records()
        
        return {
            'total': len(all_records),
            'pending': len(pending),
            'high_priority': len(high_priority),
            'file_path': str(self.file_path) if OPENPYXL_AVAILABLE else str(self.csv_path)
        }
    
    def search_records(self, keyword: str) -> List[Dict]:
        """Anahtar kelime ile kayıt ara"""
        all_records = self.get_all_records()
        keyword_lower = keyword.lower()
        
        results = []
        for record in all_records:
            for key, value in record.items():
                if keyword_lower in str(value).lower():
                    results.append(record)
                    break
        
        return results

# Global Excel instance
excel = ExcelHandler()
