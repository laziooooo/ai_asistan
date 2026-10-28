"""Excel Disa Aktarma Modulu - SQLite'dan Excel'e"""
import logging
from datetime import datetime
from pathlib import Path
from typing import List, Dict

logger = logging.getLogger(__name__)

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logger.warning("openpyxl yok - Excel disa aktarma calismaz")

from .config import DATA_DIR
from .database import db

class ExcelExporter:
    """SQLite verisini Excel'e aktar"""
    
    def __init__(self):
        self.headers = ['ID', 'Zaman', 'Gonderen', 'Kategori', 'Oncelik', 'Mesaj', 'Durum']
        self.header_colors = {
            'ID': '3B82F6',
            'Zaman': '22C55E',
            'Gonderen': 'F59E0B',
            'Kategori': '8B5CF6',
            'Oncelik': 'EF4444',
            'Mesaj': '06B6D4',
            'Durum': '6B7280'
        }
    
    def kayitlari_aktar(self, output_path: str = None) -> str:
        """Kayitlari Excel'e aktar"""
        if not OPENPYXL_AVAILABLE:
            logger.error("openpyxl yok!")
            return None
        
        try:
            # Kayitlari al
            kayitlar = db.tum_kayitlar(limit=10000)
            
            if not output_path:
                timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                output_path = DATA_DIR / f'kayitlar_{timestamp}.xlsx'
            
            # Workbook olustur
            wb = Workbook()
            ws = wb.active
            ws.title = "Kayitlar"
            
            # Basliklari yaz
            for col, header in enumerate(self.headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True, color='FFFFFF', size=11)
                color = self.header_colors.get(header, '3B82F6')
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
                cell.alignment = Alignment(horizontal='center')
            
            # Verileri yaz
            for row_idx, kayit in enumerate(kayitlar, 2):
                ws.cell(row=row_idx, column=1, value=kayit.get('id', ''))
                ws.cell(row=row_idx, column=2, value=kayit.get('zaman', ''))
                ws.cell(row=row_idx, column=3, value=kayit.get('gonderen', ''))
                ws.cell(row=row_idx, column=4, value=kayit.get('kategori', ''))
                ws.cell(row=row_idx, column=5, value=kayit.get('oncelik', ''))
                ws.cell(row=row_idx, column=6, value=kayit.get('mesaj', ''))
                ws.cell(row=row_idx, column=7, value=kayit.get('durum', ''))
            
            # Sutun genisliklerini ayarla
            ws.column_dimensions['A'].width = 8   # ID
            ws.column_dimensions['B'].width = 20  # Zaman
            ws.column_dimensions['C'].width = 15  # Gonderen
            ws.column_dimensions['D'].width = 15  # Kategori
            ws.column_dimensions['E'].width = 12  # Oncelik
            ws.column_dimensions['F'].width = 50  # Mesaj
            ws.column_dimensions['G'].width = 12  # Durum
            
            # Kaydet
            wb.save(str(output_path))
            logger.info(f"Excel'e aktarildi: {output_path} ({len(kayitlar)} kayit)")
            
            return str(output_path)
            
        except Exception as e:
            logger.error(f"Excel aktarma hatasi: {e}")
            return None
    
    def musterileri_aktar(self, output_path: str = None) -> str:
        """Musterileri Excel'e aktar"""
        if not OPENPYXL_AVAILABLE:
            logger.error("openpyxl yok!")
            return None
        
        try:
            musteriler = db.tum_musteriler()
            
            if not output_path:
                timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                output_path = DATA_DIR / f'musteriler_{timestamp}.xlsx'
            
            wb = Workbook()
            ws = wb.active
            ws.title = "Musteriler"
            
            headers = ['ID', 'Telegram ID', 'Ad', 'Telefon', 'Email', 'Mesaj Sayisi', 'Son Mesaj']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True, color='FFFFFF', size=11)
                cell.fill = PatternFill(start_color='3B82F6', end_color='3B82F6', fill_type='solid')
                cell.alignment = Alignment(horizontal='center')
            
            for row_idx, musteri in enumerate(musteriler, 2):
                ws.cell(row=row_idx, column=1, value=musteri.get('id', ''))
                ws.cell(row=row_idx, column=2, value=musteri.get('telegram_id', ''))
                ws.cell(row=row_idx, column=3, value=musteri.get('ad', ''))
                ws.cell(row=row_idx, column=4, value=musteri.get('telefon', ''))
                ws.cell(row=row_idx, column=5, value=musteri.get('email', ''))
                ws.cell(row=row_idx, column=6, value=musteri.get('toplam_mesaj', 0))
                ws.cell(row=row_idx, column=7, value=musteri.get('son_mesaj_zamani', ''))
            
            # Sutun genislikleri
            ws.column_dimensions['A'].width = 8
            ws.column_dimensions['B'].width = 15
            ws.column_dimensions['C'].width = 20
            ws.column_dimensions['D'].width = 15
            ws.column_dimensions['E'].width = 25
            ws.column_dimensions['F'].width = 12
            ws.column_dimensions['G'].width = 20
            
            wb.save(str(output_path))
            logger.info(f"Musteriler Excel'e aktarildi: {output_path}")
            
            return str(output_path)
            
        except Exception as e:
            logger.error(f"Musteri aktarma hatasi: {e}")
            return None
    
    def rapor_olustur(self, output_path: str = None) -> str:
        """Ozet rapor olustur"""
        if not OPENPYXL_AVAILABLE:
            return None
        
        try:
            istatistik = db.istatistikler()
            kategoriler = db.kategori_istatistik()
            
            if not output_path:
                timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                output_path = DATA_DIR / f'rapor_{timestamp}.xlsx'
            
            wb = Workbook()
            
            # Ozet sayfasi
            ws = wb.active
            ws.title = "Ozet"
            
            ws.cell(row=1, column=1, value="ASISTAN RAPORU").font = Font(bold=True, size=16)
            ws.cell(row=2, column=1, value=f"Tarih: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
            
            ws.cell(row=4, column=1, value="GENEL ISTATISTIKLER").font = Font(bold=True, size=12)
            
            stats = [
                ("Toplam Kayit", istatistik.get('toplam', 0)),
                ("Bekleyen", istatistik.get('bekleyen', 0)),
                ("Oncelikli", istatistik.get('oncelikli', 0)),
                ("Bugun Eklenen", istatistik.get('bugun', 0)),
            ]
            
            for idx, (label, deger) in enumerate(stats, 5):
                ws.cell(row=idx, column=1, value=label)
                ws.cell(row=idx, column=2, value=deger).font = Font(bold=True)
            
            # Kategoriler
            ws.cell(row=10, column=1, value="KATEGORI DAGILIMI").font = Font(bold=True, size=12)
            
            for idx, (kategori, sayi) in enumerate(kategoriler.items(), 11):
                ws.cell(row=idx, column=1, value=kategori)
                ws.cell(row=idx, column=2, value=sayi)
            
            ws.column_dimensions['A'].width = 25
            ws.column_dimensions['B'].width = 15
            
            wb.save(str(output_path))
            logger.info(f"Rapor olusturuldu: {output_path}")
            
            return str(output_path)
            
        except Exception as e:
            logger.error(f"Rapor olusturma hatasi: {e}")
            return None

# Global exporter instance
exporter = ExcelExporter()
